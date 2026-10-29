import pdfplumber
import re
import json
import uuid
import openpyxl
from datetime import datetime, date
from dateutil.relativedelta import relativedelta
from decimal import Decimal, InvalidOperation

from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth import update_session_auth_hash
from django.db.models import Sum, Count
from django.db import models
from django.http import HttpResponse

from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.utils import get_column_letter

from .models import (
    Pessoa, Categoria, FormaPagamento, Mes, Receita, Gasto,
    Cartao, Divida, ReservaEmergencia, Perfil
)
from .forms import (
    PessoaForm, CategoriaForm, FormaPagamentoForm, MesForm,
    ReceitaForm, GastoForm, CartaoForm, DividaForm, ReservaForm,
    PerfilForm, AlterarSenhaForm
)


# ==========================================
# PESSOAS
# ==========================================
@login_required
def pessoa_lista(request):
    pessoas = Pessoa.objects.all()
    form = PessoaForm()
    return render(request, 'financas/pessoa_lista.html', {'pessoas': pessoas, 'form': form})


@login_required
def pessoa_criar(request):
    if request.method == 'POST':
        form = PessoaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('pessoa_lista')
    else:
        form = PessoaForm()
    return render(request, 'financas/pessoa_lista.html', {
        'pessoas': Pessoa.objects.all(),
        'form': form,
    })


# ==========================================
# CATEGORIAS
# ==========================================
@login_required
def categoria_lista(request):
    categorias = Categoria.objects.all()
    form = CategoriaForm()
    return render(request, 'financas/categoria_lista.html', {'categorias': categorias, 'form': form})


@login_required
def categoria_criar(request):
    if request.method == 'POST':
        form = CategoriaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('categoria_lista')
    else:
        form = CategoriaForm()
    return render(request, 'financas/categoria_lista.html', {
        'categorias': Categoria.objects.all(),
        'form': form,
    })


@login_required
def categoria_editar(request, categoria_id):
    categoria = Categoria.objects.get(id=categoria_id)
    if request.method == 'POST':
        form = CategoriaForm(request.POST, instance=categoria)
        if form.is_valid():
            form.save()
            messages.success(request, 'Categoria atualizada com sucesso!')
            return redirect('categoria_lista')
    else:
        form = CategoriaForm(instance=categoria)
    return render(request, 'financas/categoria_form.html', {'form': form, 'categoria': categoria})


@login_required
def categoria_excluir(request, categoria_id):
    categoria = Categoria.objects.get(id=categoria_id)
    if request.method == 'POST':
        try:
            nome = categoria.nome
            categoria.delete()
            messages.success(request, f'Categoria "{nome}" excluída com sucesso!')
        except Exception:
            messages.error(request, f'Não é possível excluir "{categoria.nome}" — ela tem gastos vinculados.')
    return redirect('categoria_lista')


# ==========================================
# FORMAS DE PAGAMENTO
# ==========================================
@login_required
def formapagamento_lista(request):
    formas = FormaPagamento.objects.all()
    form = FormaPagamentoForm()
    return render(request, 'financas/formapagamento_lista.html', {'formas': formas, 'form': form})


@login_required
def formapagamento_criar(request):
    if request.method == 'POST':
        form = FormaPagamentoForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('formapagamento_lista')
    else:
        form = FormaPagamentoForm()
    return render(request, 'financas/formapagamento_lista.html', {
        'formas': FormaPagamento.objects.all(),
        'form': form,
    })


# ==========================================
# MESES
# ==========================================
@login_required
def mes_lista(request):
    meses_abertos = Mes.objects.filter(fechado=False).order_by('-ano', '-mes')
    meses_fechados = Mes.objects.filter(fechado=True).order_by('-ano', '-mes')
    form = MesForm()
    mes_ativo_id = request.session.get('mes_ativo_id')

    return render(request, 'financas/mes_lista.html', {
        'meses': meses_abertos,
        'meses_fechados': meses_fechados,
        'form': form,
        'mes_ativo_id': mes_ativo_id,
    })


@login_required
def mes_criar(request):
    if request.method == 'POST':
        form = MesForm(request.POST)
        if form.is_valid():
            novo_mes = form.save(commit=False)

            data_novo = date(novo_mes.ano, novo_mes.mes, 1)
            data_anterior = data_novo - relativedelta(months=1)

            mes_anterior = Mes.objects.filter(ano=data_anterior.year, mes=data_anterior.month).first()

            if not mes_anterior:
                mes_anterior = Mes.objects.filter(receitas__isnull=False).order_by('-ano', '-mes').first()

            novo_mes.save()

            if mes_anterior:
                receitas_antigas = Receita.objects.filter(mes=mes_anterior)

                if receitas_antigas.exists():
                    for r in receitas_antigas:
                        Receita.objects.create(
                            mes=novo_mes,
                            descricao=r.descricao,
                            valor=r.valor,
                            data=date(novo_mes.ano, novo_mes.mes, 1),
                            pessoa=r.pessoa
                        )
                    messages.success(request, f'Mês {novo_mes} criado e {receitas_antigas.count()} receitas copiadas de {mes_anterior}!')
                else:
                    messages.warning(request, f'Mês {novo_mes} criado, mas o mês anterior ({mes_anterior}) não tinha receitas para copiar.')
            else:
                messages.success(request, f'Mês {novo_mes} criado! Nenhum mês anterior com receitas foi encontrado.')

            return redirect('mes_lista')
    else:
        form = MesForm()

    return render(request, 'financas/mes_lista.html', {
        'meses': Mes.objects.filter(fechado=False),
        'meses_fechados': Mes.objects.filter(fechado=True),
        'form': form,
        'mes_ativo_id': request.session.get('mes_ativo_id'),
    })


@login_required
def mes_selecionar(request, mes_id):
    if request.method == 'POST':
        request.session['mes_ativo_id'] = mes_id
    return redirect('mes_lista')


@login_required
def mes_fechar(request, mes_id):
    if request.method == 'POST':
        mes_atual = Mes.objects.get(id=mes_id)
        if mes_atual.fechado:
            return redirect('mes_lista')

        mes_atual.fechado = True
        mes_atual.save()

        data_atual = date(mes_atual.ano, mes_atual.mes, 1)
        data_proximo = data_atual + relativedelta(months=1)

        proximo_mes, criado = Mes.objects.get_or_create(
            ano=data_proximo.year,
            mes=data_proximo.month,
            defaults={'fechado': False}
        )

        if criado:
            receitas_antigas = Receita.objects.filter(mes=mes_atual)
            for r in receitas_antigas:
                Receita.objects.create(
                    mes=proximo_mes,
                    descricao=r.descricao,
                    valor=r.valor,
                    data=data_proximo,
                    pessoa=r.pessoa
                )
            messages.success(request, f'Mês {mes_atual} fechado! As receitas foram copiadas para {proximo_mes}.')
        else:
            messages.success(request, f'Mês {mes_atual} fechado! O mês {proximo_mes} já estava aberto.')

        request.session['mes_ativo_id'] = proximo_mes.id

    return redirect('mes_lista')


@login_required
def mes_reabrir(request, mes_id):
    if request.method == 'POST':
        mes = Mes.objects.get(id=mes_id)
        mes.fechado = False
        mes.save()
        messages.success(request, f'Mês {mes} reaberto com sucesso!')
    return redirect('mes_lista')


@login_required
def mes_clonar(request, mes_id):
    if request.method == 'POST':
        mes_atual = Mes.objects.get(id=mes_id)
        data_atual = date(mes_atual.ano, mes_atual.mes, 1)
        data_anterior = data_atual - relativedelta(months=1)

        mes_anterior = Mes.objects.filter(ano=data_anterior.year, mes=data_anterior.month).first()

        if not mes_anterior:
            messages.error(request, 'Não encontrei o mês anterior para clonar.')
            return redirect('mes_lista')

        receitas_antigas = Receita.objects.filter(mes=mes_anterior)
        for r in receitas_antigas:
            Receita.objects.create(
                mes=mes_atual,
                descricao=r.descricao,
                valor=r.valor,
                data=date(mes_atual.ano, mes_atual.mes, 1),
                pessoa=r.pessoa
            )

        gastos_antigos = Gasto.objects.filter(mes=mes_anterior, cartao__isnull=True)
        for g in gastos_antigos:
            Gasto.objects.create(
                mes=mes_atual,
                descricao=g.descricao,
                valor=g.valor,
                data=date(mes_atual.ano, mes_atual.mes, 1),
                categoria=g.categoria,
                pessoa=g.pessoa,
                forma_pagamento=g.forma_pagamento,
                cartao=None
            )

        messages.success(request, f'Mês clonado com sucesso a partir de {mes_anterior}!')
        return redirect('mes_lista')

    return redirect('mes_lista')


# ==========================================
# RECEITAS
# ==========================================
@login_required
def receita_lista(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return render(request, 'financas/receita_lista.html', {
            'sem_mes_ativo': True,
            'receitas': [],
            'form': None,
        })

    mes_ativo = Mes.objects.get(id=mes_ativo_id)
    receitas = Receita.objects.filter(mes=mes_ativo)
    form = ReceitaForm()
    total = sum(r.valor for r in receitas)

    return render(request, 'financas/receita_lista.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'receitas': receitas,
        'form': form,
        'total': total,
    })


@login_required
def receita_criar(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return redirect('mes_lista')

    if request.method == 'POST':
        form = ReceitaForm(request.POST)
        if form.is_valid():
            receita = form.save(commit=False)
            receita.mes = Mes.objects.get(id=mes_ativo_id)
            receita.save()
            return redirect('receita_lista')
    else:
        form = ReceitaForm()

    return render(request, 'financas/receita_lista.html', {
        'sem_mes_ativo': False,
        'mes_ativo': Mes.objects.get(id=mes_ativo_id),
        'receitas': Receita.objects.filter(mes_id=mes_ativo_id),
        'form': form,
        'total': sum(r.valor for r in Receita.objects.filter(mes_id=mes_ativo_id)),
    })


@login_required
def receita_excluir(request, receita_id):
    receita = Receita.objects.get(id=receita_id)
    if request.method == 'POST':
        receita.delete()
        messages.success(request, 'Receita excluída com sucesso!')
    return redirect('receita_lista')


@login_required
def receita_editar(request, receita_id):
    receita = Receita.objects.get(id=receita_id)
    if request.method == 'POST':
        form = ReceitaForm(request.POST, instance=receita)
        if form.is_valid():
            form.save()
            messages.success(request, 'Receita atualizada com sucesso!')
            return redirect('receita_lista')
    else:
        form = ReceitaForm(instance=receita)
    return render(request, 'financas/receita_form.html', {'form': form, 'receita': receita})


# ==========================================
# GASTOS
# ==========================================
@login_required
def gasto_lista(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return render(request, 'financas/gasto_lista.html', {
            'sem_mes_ativo': True,
            'gastos': [],
            'form': None,
        })

    mes_ativo = Mes.objects.get(id=mes_ativo_id)
    gastos = Gasto.objects.filter(mes=mes_ativo)
    form = GastoForm()
    total = sum(g.valor for g in gastos)

    return render(request, 'financas/gasto_lista.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'gastos': gastos,
        'form': form,
        'total': total,
    })


@login_required
def gasto_criar(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return redirect('mes_lista')

    if request.method == 'POST':
        form = GastoForm(request.POST)
        if form.is_valid():
            gasto = form.save(commit=False)
            mes_ativo = Mes.objects.get(id=mes_ativo_id)

            parcela_atual = form.cleaned_data.get('parcela_atual') or 1
            total_parcelas = form.cleaned_data.get('total_parcelas') or 1

            if parcela_atual > total_parcelas:
                parcela_atual = total_parcelas

            if total_parcelas > 1:
                grupo_id = uuid.uuid4()
                data_base = gasto.data

                for i in range(parcela_atual, total_parcelas + 1):
                    nova_data = data_base + relativedelta(months=(i - parcela_atual))

                    if i == parcela_atual:
                        mes_parcela = mes_ativo
                    else:
                        mes_parcela, _ = Mes.objects.get_or_create(ano=nova_data.year, mes=nova_data.month)

                    novo_gasto = Gasto(
                        mes=mes_parcela,
                        descricao=f"{gasto.descricao} ({i}/{total_parcelas})",
                        valor=gasto.valor,
                        data=nova_data,
                        categoria=gasto.categoria,
                        pessoa=gasto.pessoa,
                        forma_pagamento=gasto.forma_pagamento,
                        cartao=gasto.cartao,
                        parcela_atual=i,
                        total_parcelas=total_parcelas,
                        grupo_parcelas=grupo_id
                    )
                    novo_gasto.save()

                messages.success(request, f"Gasto parcelado em {total_parcelas}x com sucesso!")
                return redirect('gasto_lista')

            else:
                gasto.mes = mes_ativo
                gasto.parcela_atual = 1
                gasto.total_parcelas = 1
                gasto.save()
                messages.success(request, "Gasto cadastrado com sucesso!")
                return redirect('gasto_lista')
    else:
        form = GastoForm()

    mes_ativo = Mes.objects.get(id=mes_ativo_id)
    gastos = Gasto.objects.filter(mes=mes_ativo)

    return render(request, 'financas/gasto_lista.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'gastos': gastos,
        'form': form,
        'total': sum(g.valor for g in gastos),
    })


@login_required
def gasto_excluir(request, gasto_id):
    gasto = Gasto.objects.get(id=gasto_id)
    if request.method == 'POST':
        gasto.delete()
        messages.success(request, 'Gasto excluído com sucesso!')
    return redirect('gasto_lista')


@login_required
def gasto_editar(request, gasto_id):
    gasto = Gasto.objects.get(id=gasto_id)
    if request.method == 'POST':
        form = GastoForm(request.POST, instance=gasto)
        if form.is_valid():
            form.save()
            messages.success(request, 'Gasto atualizado com sucesso!')
            return redirect('gasto_lista')
    else:
        form = GastoForm(instance=gasto)
    return render(request, 'financas/gasto_form.html', {'form': form, 'gasto': gasto})


# ==========================================
# CARTÕES
# ==========================================
@login_required
def cartao_lista(request):
    cartoes = Cartao.objects.all()
    form = CartaoForm()
    return render(request, 'financas/cartao_lista.html', {'cartoes': cartoes, 'form': form})


@login_required
def cartao_criar(request):
    if request.method == 'POST':
        form = CartaoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cartão cadastrado com sucesso!')
            return redirect('cartao_lista')
    else:
        form = CartaoForm()
    return render(request, 'financas/cartao_lista.html', {
        'cartoes': Cartao.objects.all(),
        'form': form,
    })


@login_required
def cartao_editar(request, cartao_id):
    cartao = Cartao.objects.get(id=cartao_id)
    if request.method == 'POST':
        form = CartaoForm(request.POST, instance=cartao)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cartão atualizado com sucesso!')
            return redirect('cartao_lista')
    else:
        form = CartaoForm(instance=cartao)
    return render(request, 'financas/cartao_form.html', {'form': form, 'cartao': cartao})


@login_required
def cartao_excluir(request, cartao_id):
    cartao = Cartao.objects.get(id=cartao_id)
    if request.method == 'POST':
        try:
            nome = cartao.nome
            cartao.delete()
            messages.success(request, f'Cartão "{nome}" excluído com sucesso!')
        except Exception:
            messages.error(request, f'Não é possível excluir "{cartao.nome}" — ele tem gastos vinculados.')
    return redirect('cartao_lista')


# ==========================================
# DASHBOARD
# ==========================================
@login_required
def dashboard(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return render(request, 'financas/dashboard.html', {'sem_mes_ativo': True})

    mes_ativo = Mes.objects.get(id=mes_ativo_id)

    total_receitas = Receita.objects.filter(mes=mes_ativo).aggregate(total=Sum('valor'))['total'] or Decimal('0')
    total_gastos = Gasto.objects.filter(mes=mes_ativo).aggregate(total=Sum('valor'))['total'] or Decimal('0')
    verba = total_receitas - total_gastos

    if total_receitas == 0:
        alerta = 'sem_receita'
    elif verba < 0:
        alerta = 'estourado'
    elif verba < (total_receitas * Decimal('0.10')):
        alerta = 'apertado'
    else:
        alerta = 'ok'

    gastos_essenciais = Gasto.objects.filter(
        mes=mes_ativo, categoria__tipo='E'
    ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

    gastos_nao_essenciais = Gasto.objects.filter(
        mes=mes_ativo, categoria__tipo='N'
    ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

    if total_receitas > 0:
        pct_essenciais = float((gastos_essenciais / total_receitas) * 100)
        pct_nao_essenciais = float((gastos_nao_essenciais / total_receitas) * 100)
        pct_sobra = float((verba / total_receitas) * 100)
    else:
        pct_essenciais = 0
        pct_nao_essenciais = 0
        pct_sobra = 0

    categorias_com_limite = Categoria.objects.filter(
        ativo=True, limite_mensal__isnull=False
    ).exclude(limite_mensal=0)

    orcamentos = []
    for cat in categorias_com_limite:
        gasto_cat = Gasto.objects.filter(
            mes=mes_ativo, categoria=cat
        ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

        pct_usado = float((gasto_cat / cat.limite_mensal) * 100) if cat.limite_mensal > 0 else 0
        restante = cat.limite_mensal - gasto_cat

        if pct_usado > 100:
            cor = 'danger'
        elif pct_usado >= 70:
            cor = 'warning'
        else:
            cor = 'success'

        pct_usado_exibicao = min(pct_usado, 100)

        orcamentos.append({
            'categoria': cat,
            'gasto': gasto_cat,
            'limite': cat.limite_mensal,
            'restante': restante,
            'pct_usado': pct_usado,
            'pct_usado_exibicao': pct_usado_exibicao,
            'cor': cor,
        })

    cartoes_ativos = Cartao.objects.filter(ativo=True)
    resumo_cartoes = []

    for c in cartoes_ativos:
        total_cartao = Gasto.objects.filter(
            mes=mes_ativo, cartao=c
        ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

        resumo_cartoes.append({
            'cartao': c,
            'total_mes': total_cartao
        })

    ultimas_receitas = Receita.objects.filter(mes=mes_ativo).order_by('-id')[:5]
    ultimos_gastos = Gasto.objects.filter(mes=mes_ativo).order_by('-id')[:5]

    gastos_por_categoria = Gasto.objects.filter(mes=mes_ativo).values('categoria__nome').annotate(
        total=Sum('valor')
    ).order_by('-total')

    grafico_labels = [item['categoria__nome'] for item in gastos_por_categoria]
    grafico_valores = [float(item['total']) for item in gastos_por_categoria]

    return render(request, 'financas/dashboard.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'total_receitas': total_receitas,
        'total_gastos': total_gastos,
        'verba': verba,
        'alerta': alerta,
        'gastos_essenciais': gastos_essenciais,
        'gastos_nao_essenciais': gastos_nao_essenciais,
        'pct_essenciais': pct_essenciais,
        'pct_nao_essenciais': pct_nao_essenciais,
        'pct_sobra': pct_sobra,
        'orcamentos': orcamentos,
        'resumo_cartoes': resumo_cartoes,
        'ultimas_receitas': ultimas_receitas,
        'ultimos_gastos': ultimos_gastos,
        'grafico_labels': grafico_labels,
        'grafico_valores': grafico_valores,
    })


# ==========================================
# IMPORTAÇÃO DE FATURA (PDF)
# ==========================================
@login_required
def importar_fatura(request):
    # 1. SALVAR NO BANCO
    if request.method == 'POST' and 'salvar_importacao' in request.POST:
        mes_ativo_id = request.session.get('mes_ativo_id')
        if not mes_ativo_id:
            messages.error(request, 'Nenhum mês ativo selecionado.')
            return redirect('mes_lista')

        mes_ativo = Mes.objects.get(id=mes_ativo_id)
        categoria_id = request.POST.get('categoria')
        pessoa_id = request.POST.get('pessoa')
        cartao_id = request.POST.get('cartao')

        datas = request.POST.getlist('data[]')
        descricoes = request.POST.getlist('descricao[]')
        valores = request.POST.getlist('valor[]')
        parcelas = request.POST.getlist('parcela[]')

        if not (categoria_id and pessoa_id and cartao_id and datas):
            messages.error(request, 'Preencha todos os campos e selecione pelo menos um gasto.')
            return redirect('importar_fatura')

        categoria = Categoria.objects.get(id=categoria_id)
        pessoa = Pessoa.objects.get(id=pessoa_id)
        cartao = Cartao.objects.get(id=cartao_id)

        forma_credito, _ = FormaPagamento.objects.get_or_create(nome="Cartão de Crédito")
        salvos = 0

        for i in range(len(datas)):
            descricao_original = descricoes[i]
            descricao_limpa = re.sub(r'^[•\s\d]+', '', descricao_original).strip()

            try:
                valor = Decimal(str(valores[i]))
            except InvalidOperation:
                continue

            try:
                data_gasto = datetime.strptime(datas[i], '%Y-%m-%d').date()
            except ValueError:
                continue

            parcela_info = parcelas[i] if i < len(parcelas) else ""
            match_parcela = re.search(r'(\d+)/(\d+)', parcela_info)

            if not match_parcela:
                match_parcela = re.search(r'Parcela (\d+)/(\d+)', descricao_limpa)

            if match_parcela:
                parcela_atual = int(match_parcela.group(1))
                total_parcelas = int(match_parcela.group(2))
                parcelas_restantes = total_parcelas - parcela_atual

                nome_base = re.sub(r'Parcela \d+/\d+', '', descricao_limpa).strip()
                nome_base = re.sub(r'\d{2}/\d{2}$', '', nome_base).strip()

                Gasto.objects.create(
                    mes=mes_ativo,
                    descricao=f"{nome_base} ({parcela_atual}/{total_parcelas})",
                    valor=valor,
                    data=data_gasto,
                    categoria=categoria,
                    pessoa=pessoa,
                    forma_pagamento=forma_credito,
                    cartao=cartao,
                )

                for j in range(1, parcelas_restantes + 1):
                    mes_futuro_num = mes_ativo.mes + j
                    ano_futuro = mes_ativo.ano
                    if mes_futuro_num > 12:
                        mes_futuro_num -= 12
                        ano_futuro += 1

                    mes_futuro, _ = Mes.objects.get_or_create(ano=ano_futuro, mes=mes_futuro_num)

                    nova_data = data_gasto.replace(year=ano_futuro, month=mes_futuro_num)
                    Gasto.objects.create(
                        mes=mes_futuro,
                        descricao=f"{nome_base} ({parcela_atual + j}/{total_parcelas})",
                        valor=valor,
                        data=nova_data,
                        categoria=categoria,
                        pessoa=pessoa,
                        forma_pagamento=forma_credito,
                        cartao=cartao,
                    )
            else:
                Gasto.objects.create(
                    mes=mes_ativo,
                    descricao=descricao_limpa,
                    valor=valor,
                    data=data_gasto,
                    categoria=categoria,
                    pessoa=pessoa,
                    forma_pagamento=forma_credito,
                    cartao=cartao,
                )
            salvos += 1

        messages.success(request, f'{salvos} transações importadas com sucesso!')
        return redirect('gasto_lista')

    # 2. LER O PDF
    if request.method == 'POST' and request.FILES.get('arquivo_pdf'):
        mes_ativo_id = request.session.get('mes_ativo_id')
        if not mes_ativo_id:
            messages.error(request, 'Nenhum mês ativo selecionado.')
            return redirect('mes_lista')
        mes_ativo = Mes.objects.get(id=mes_ativo_id)

        pdf_file = request.FILES.get('arquivo_pdf')
        transacoes = []

        try:
            with pdfplumber.open(pdf_file) as pdf:
                texto_completo = "".join([page.extract_text() or "" for page in pdf.pages])
                is_santander = "Detalhamento da Fatura" in texto_completo

                for page in pdf.pages:
                    texto = page.extract_text()
                    if not texto:
                        continue

                    linhas = [l.strip() for l in texto.splitlines() if l.strip()]

                    if is_santander:
                        secao_atual = None
                        dentro_detalhamento = False

                        padrao_item = re.compile(
                            r'^(?:\d+\s+)?'
                            r'(\d{2}/\d{2})\s*'
                            r'(.+?)\s*'
                            r'(?:(\d{2}/\d{2})\s*)?'
                            r'([\d\.]+,\d{2})'
                        )

                        for linha in linhas:
                            if "Detalhamento da Fatura" in linha:
                                dentro_detalhamento = True
                                continue
                            if not dentro_detalhamento:
                                continue

                            linha_lower = linha.lower()

                            if "pagamento e demais" in linha_lower or "pagamento de fatura" in linha_lower:
                                secao_atual = "ignorar"
                                continue

                            if linha_lower.startswith("parcelamentos"):
                                secao_atual = "parcelamentos"
                                continue
                            if linha_lower.startswith("despesas"):
                                secao_atual = "despesas"
                                continue
                            if linha_lower.startswith("valor total"):
                                secao_atual = None
                                continue

                            if secao_atual in ["parcelamentos", "despesas"]:
                                m = padrao_item.match(linha)
                                if m:
                                    data, descricao, parcela, valor_brl = m.groups()

                                    try:
                                        valor_limpo = valor_brl.replace('.', '').replace(',', '.')
                                        valor = Decimal(valor_limpo)
                                    except InvalidOperation:
                                        continue

                                    if valor <= 0:
                                        valor = abs(valor)

                                    dia, mes = data.split('/')
                                    ano = mes_ativo.ano
                                    if int(mes) > mes_ativo.mes:
                                        ano = ano - 1

                                    try:
                                        data_gasto = date(ano, int(mes), int(dia))
                                    except ValueError:
                                        continue

                                    descricao_limpa = re.sub(r'^[•\s\d]+', '', descricao).strip()
                                    descricao_limpa = re.sub(r'\d{2}/\d{2}$', '', descricao_limpa).strip()

                                    if not descricao_limpa:
                                        continue

                                    transacoes.append({
                                        'data': data_gasto.isoformat(),
                                        'descricao': descricao_limpa,
                                        'parcela': parcela if parcela else "",
                                        'valor': str(valor),
                                    })
                                continue

                    else:
                        for linha in linhas:
                            if " a " in linha or "Total" in linha or "Limite" in linha:
                                continue

                            match_nubank = re.search(
                                r'(\d{2}\s[A-Z]{3})\s*(?:\d{4,6})?\s*(.*?)\s*R\$\s*([\d\.]+,\d{2})',
                                linha
                            )
                            if match_nubank:
                                data_str = match_nubank.group(1)
                                descricao = match_nubank.group(2).strip()
                                valor_str = match_nubank.group(3)

                                if not descricao or descricao.startswith("Pagamento em"):
                                    continue

                                try:
                                    valor_limpo = valor_str.replace('.', '').replace(',', '.')
                                    valor = Decimal(valor_limpo)
                                except InvalidOperation:
                                    continue

                                meses = {
                                    'JAN': 1, 'FEV': 2, 'MAR': 3, 'ABR': 4,
                                    'MAI': 5, 'JUN': 6, 'JUL': 7, 'AGO': 8,
                                    'SET': 9, 'OUT': 10, 'NOV': 11, 'DEZ': 12,
                                }
                                try:
                                    dia, mes_abrev = data_str.split()
                                    mes = meses.get(mes_abrev, 1)
                                except ValueError:
                                    continue

                                ano = datetime.now().year
                                if mes > mes_ativo.mes:
                                    ano = ano - 1

                                try:
                                    data_gasto = date(ano, mes, int(dia))
                                except ValueError:
                                    continue

                                descricao_limpa = re.sub(r'^[•\s\d]+', '', descricao).strip()
                                if not descricao_limpa:
                                    continue

                                transacoes.append({
                                    'data': data_gasto.isoformat(),
                                    'descricao': descricao_limpa,
                                    'parcela': '',
                                    'valor': str(valor),
                                })

        except Exception as e:
            messages.error(request, f'Erro ao ler PDF: {e}')
            return redirect('importar_fatura')

        if not transacoes:
            messages.warning(request, 'Nenhuma transação válida encontrada no PDF.')

        categorias = Categoria.objects.filter(ativo=True)
        pessoas = Pessoa.objects.filter(ativo=True)
        cartoes = Cartao.objects.filter(ativo=True)

        return render(request, 'financas/importar_fatura.html', {
            'transacoes': transacoes,
            'total': sum(Decimal(t['valor']) for t in transacoes),
            'categorias': categorias,
            'pessoas': pessoas,
            'cartoes': cartoes,
        })

    return render(request, 'financas/importar_fatura.html')


# ==========================================
# PROJEÇÃO
# ==========================================
@login_required
def projecao(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return render(request, 'financas/projecao.html', {'sem_mes_ativo': True})

    mes_ativo = Mes.objects.get(id=mes_ativo_id)

    ano_base = int(request.GET.get('ano', mes_ativo.ano))
    mes_base = int(request.GET.get('mes', mes_ativo.mes))

    nomes_meses = {
        1: 'Janeiro', 2: 'Fevereiro', 3: 'Março', 4: 'Abril', 5: 'Maio', 6: 'Junho',
        7: 'Julho', 8: 'Agosto', 9: 'Setembro', 10: 'Outubro', 11: 'Novembro', 12: 'Dezembro'
    }

    meses_colunas = []
    data_atual = date(ano_base, mes_base, 1)

    for i in range(12):
        mes_alvo = data_atual + relativedelta(months=i)
        meses_colunas.append({
            'id': f"{mes_alvo.year}-{mes_alvo.month}",
            'nome': f"{nomes_meses[mes_alvo.month]}/{mes_alvo.year}",
            'ano': mes_alvo.year,
            'mes': mes_alvo.month,
        })

    data_anterior = data_atual - relativedelta(months=1)
    data_posterior = data_atual + relativedelta(months=1)

    anos_filtro = list(set([m['ano'] for m in meses_colunas]))

    gastos_periodo = Gasto.objects.filter(
        mes__ano__in=anos_filtro
    ).select_related('mes', 'cartao', 'pessoa')

    totais_gastos = {m['id']: Decimal('0') for m in meses_colunas}
    totais_receitas = {m['id']: Decimal('0') for m in meses_colunas}
    saldos_livres = {m['id']: Decimal('0') for m in meses_colunas}

    agrupado = {}

    for g in gastos_periodo:
        id_mes = f"{g.mes.ano}-{g.mes.mes}"

        if id_mes not in totais_gastos:
            continue

        nome_limpo = re.sub(r'\s*\(?\d+/\d+\)?|\s*-\s*Parcela\s*\d+/\d+', '', g.descricao).strip()

        if not nome_limpo:
            nome_limpo = g.descricao

        chave_grupo = f"{g.pessoa.nome}|{nome_limpo}"

        if chave_grupo not in agrupado:
            agrupado[chave_grupo] = {
                'categoria': nome_limpo,
                'pessoa': g.pessoa.nome,
                'valores': {}
            }

        if id_mes not in agrupado[chave_grupo]['valores']:
            agrupado[chave_grupo]['valores'][id_mes] = Decimal('0')

        agrupado[chave_grupo]['valores'][id_mes] += g.valor
        totais_gastos[id_mes] += g.valor

    receitas_periodo = Receita.objects.filter(mes__ano__in=anos_filtro)
    for r in receitas_periodo:
        id_mes = f"{r.mes.ano}-{r.mes.mes}"
        if id_mes in totais_receitas:
            totais_receitas[id_mes] += r.valor

    for m in meses_colunas:
        id_mes = m['id']
        saldos_livres[id_mes] = totais_receitas[id_mes] - totais_gastos[id_mes]

    linhas_projecao = list(agrupado.values())
    linhas_projecao = sorted(linhas_projecao, key=lambda x: (x['pessoa'], x['categoria'].lower()))

    return render(request, 'financas/projecao.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'meses_colunas': meses_colunas,
        'linhas_projecao': linhas_projecao,
        'totais_gastos': totais_gastos,
        'totais_receitas': totais_receitas,
        'saldos_livres': saldos_livres,
        'ano_base': ano_base,
        'mes_base': mes_base,
        'mes_base_nome': f"{nomes_meses[mes_base]}/{ano_base}",
        'ano_anterior': data_anterior.year,
        'mes_anterior': data_anterior.month,
        'ano_posterior': data_posterior.year,
        'mes_posterior': data_posterior.month,
    })


# ==========================================
# HISTÓRICO
# ==========================================
@login_required
def historico_lista(request):
    meses_fechados = Mes.objects.filter(fechado=True).order_by('-ano', '-mes')
    return render(request, 'financas/historico_lista.html', {'meses': meses_fechados})


@login_required
def historico_detalhe(request, mes_id):
    mes = Mes.objects.get(id=mes_id)
    receitas = Receita.objects.filter(mes=mes)
    gastos = Gasto.objects.filter(mes=mes)

    total_receitas = receitas.aggregate(total=Sum('valor'))['total'] or Decimal('0')
    total_gastos = gastos.aggregate(total=Sum('valor'))['total'] or Decimal('0')
    sobra = total_receitas - total_gastos

    return render(request, 'financas/historico_detalhe.html', {
        'mes': mes,
        'receitas': receitas,
        'gastos': gastos,
        'total_receitas': total_receitas,
        'total_gastos': total_gastos,
        'sobra': sobra,
    })


# ==========================================
# COFRE (DÍVIDAS E RESERVA)
# ==========================================
@login_required
def cofre(request):
    dividas = Divida.objects.filter(ativo=True)
    reserva = ReservaEmergencia.objects.first()

    if not reserva:
        reserva = ReservaEmergencia.objects.create()

    form_divida = DividaForm()
    form_reserva = ReservaForm(instance=reserva)

    total_gastos_historico = Gasto.objects.aggregate(total=Sum('valor'))['total'] or Decimal('0')
    meses_registrados = Mes.objects.count()
    custo_vida_medio = total_gastos_historico / meses_registrados if meses_registrados > 0 else Decimal('0')

    meta_reserva = custo_vida_medio * reserva.meses_alvo
    pct_reserva = (reserva.valor_guardado / meta_reserva * 100) if meta_reserva > 0 else 0
    total_divida = sum(d.falta_pagar for d in dividas)

    if request.method == 'POST':
        if 'salvar_divida' in request.POST:
            form_divida = DividaForm(request.POST)
            if form_divida.is_valid():
                form_divida.save()
                messages.success(request, 'Dívida cadastrada!')
                return redirect('cofre')

        elif 'salvar_reserva' in request.POST:
            form_reserva = ReservaForm(request.POST, instance=reserva)
            if form_reserva.is_valid():
                form_reserva.save()
                messages.success(request, 'Reserva atualizada!')
                return redirect('cofre')

    return render(request, 'financas/cofre.html', {
        'dividas': dividas,
        'reserva': reserva,
        'form_divida': form_divida,
        'form_reserva': form_reserva,
        'total_divida': total_divida,
        'custo_vida_medio': custo_vida_medio,
        'meta_reserva': meta_reserva,
        'pct_reserva': float(pct_reserva),
    })


@login_required
def divida_excluir(request, divida_id):
    divida = Divida.objects.get(id=divida_id)
    if request.method == 'POST':
        divida.delete()
        messages.success(request, 'Dívida removida!')
    return redirect('cofre')


# ==========================================
# RELATÓRIOS
# ==========================================
@login_required
def relatorios(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return render(request, 'financas/relatorios.html', {'sem_mes_ativo': True})

    mes_ativo = Mes.objects.get(id=mes_ativo_id)

    gastos_por_pessoa = Gasto.objects.filter(mes=mes_ativo).values('pessoa__nome').annotate(
        total=Sum('valor')
    ).order_by('-total')

    gastos_mes = Gasto.objects.filter(mes=mes_ativo)
    total_gasto_mes = sum(g.valor for g in gastos_mes)

    agrupado_mes = {}
    for g in gastos_mes:
        nome_limpo = re.sub(r'\s*\(\d+/\d+\)|\s*-\s*Parcela\s*\d+/\d+', '', g.descricao).strip()

        if g.cartao:
            chave = f"{nome_limpo} ({g.cartao.nome})"
            tipo_visual = 'Cartão'
        else:
            chave = g.categoria.nome
            tipo_visual = 'Essencial' if g.categoria.tipo == 'E' else 'Não Essencial'

        if chave not in agrupado_mes:
            agrupado_mes[chave] = {'nome': chave, 'tipo': tipo_visual, 'total': Decimal('0')}

        agrupado_mes[chave]['total'] += g.valor

    top_categorias = sorted(agrupado_mes.values(), key=lambda x: x['total'], reverse=True)[:5]

    ultimos_meses = Mes.objects.filter(gastos__isnull=False).order_by('-ano', '-mes')[:3]
    media_categorias = []
    if ultimos_meses:
        ids_meses = [m.id for m in ultimos_meses]
        gastos_agrupados = Gasto.objects.filter(mes__id__in=ids_meses).values(
            'categoria__nome'
        ).annotate(
            total=Sum('valor')
        ).order_by('-total')[:5]

        for item in gastos_agrupados:
            media = item['total'] / 3
            media_categorias.append({
                'nome': item['categoria__nome'],
                'media': media
            })

    lista_completa_gastos = Gasto.objects.filter(mes=mes_ativo).order_by('-data', '-id')

    dicas = []
    total_receitas = Receita.objects.filter(mes=mes_ativo).aggregate(total=Sum('valor'))['total'] or Decimal('0')
    saldo = total_receitas - total_gasto_mes

    if top_categorias and top_categorias[0]['total'] > (total_gasto_mes * Decimal('0.3')):
        dicas.append(f"Você gastou mais de 30% do seu dinheiro em '{top_categorias[0]['nome']}'. Tente negociar ou evitar novas compras nesse local.")

    if total_gasto_mes > total_receitas and total_receitas > 0:
        dicas.append("Atenção! Seus gastos passaram das suas receitas este mês. Reveja os gastos não essenciais.")
    elif total_receitas == 0:
        dicas.append("Você ainda não lançou suas receitas este mês. Vá na aba Receitas e lance seu salário para ver o saldo real.")
    else:
        if saldo > 0:
            dicas.append("Parabéns! Suas contas parecem equilibradas este mês. Que tal guardar a sobra na sua Reserva de Emergência?")

    categorias_estouradas = Categoria.objects.filter(ativo=True, limite_mensal__isnull=False)
    for cat in categorias_estouradas:
        gasto_cat = Gasto.objects.filter(mes=mes_ativo, categoria=cat).aggregate(total=Sum('valor'))['total'] or Decimal('0')
        if gasto_cat > cat.limite_mensal:
            dicas.append(f"O orçamento de '{cat.nome}' estourou! Você gastou R$ {gasto_cat} mas o limite era R$ {cat.limite_mensal}.")

    dicas = dicas[:3]

    return render(request, 'financas/relatorios.html', {
        'sem_mes_ativo': False,
        'mes_ativo': mes_ativo,
        'gastos_por_pessoa': gastos_por_pessoa,
        'top_categorias': top_categorias,
        'media_categorias': media_categorias,
        'lista_completa_gastos': lista_completa_gastos,
        'total_gasto_mes': total_gasto_mes,
        'dicas': dicas,
    })


# ==========================================
# EXPORTAR RELATÓRIO EXCEL
# ==========================================
@login_required
def exportar_relatorio(request):
    mes_ativo_id = request.session.get('mes_ativo_id')
    if not mes_ativo_id:
        return redirect('mes_lista')

    mes_ativo = Mes.objects.get(id=mes_ativo_id)

    gastos = Gasto.objects.filter(mes=mes_ativo).order_by('-data')
    total_gasto_mes = sum(g.valor for g in gastos)

    gastos_por_categoria = Gasto.objects.filter(mes=mes_ativo).values('categoria__nome').annotate(
        total=Sum('valor')
    ).order_by('-total')

    meses_futuros = Mes.objects.filter(
        (models.Q(ano__gt=mes_ativo.ano) | models.Q(ano=mes_ativo.ano, mes__gt=mes_ativo.mes))
    ).order_by('ano', 'mes')[:6]

    projecao_dados = []
    for m in meses_futuros:
        g_total = Gasto.objects.filter(mes=m).aggregate(total=Sum('valor'))['total'] or Decimal('0')
        r_total = Receita.objects.filter(mes=m).aggregate(total=Sum('valor'))['total'] or Decimal('0')
        sobra = r_total - g_total
        projecao_dados.append({
            'nome': f"{m.mes}/{m.ano}",
            'gasto': g_total,
            'receita': r_total,
            'sobra': sobra
        })

    wb = openpyxl.Workbook()

    # ABA 1: DASHBOARD
    ws1 = wb.active
    ws1.title = "Dashboard"

    ws1['A1'] = f"RELATÓRIO FINANCEIRO - {mes_ativo}"
    ws1['A1'].font = Font(name='Calibri', bold=True, size=16, color='1E1B4B')
    ws1.merge_cells('A1:F1')

    ws1['E2'] = "Categoria"
    ws1['F2'] = "Valor"
    for i, item in enumerate(gastos_por_categoria, start=3):
        ws1[f'E{i}'] = item['categoria__nome']
        ws1[f'F{i}'] = float(item['total'])

    pie = PieChart()
    labels = Reference(ws1, min_col=5, min_row=3, max_row=2 + len(gastos_por_categoria))
    data = Reference(ws1, min_col=6, min_row=2, max_row=2 + len(gastos_por_categoria))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = "Divisão de Gastos por Categoria"
    ws1.add_chart(pie, "A3")

    # ABA 2: PROJEÇÃO
    ws2 = wb.create_sheet(title="Projeção 6 Meses")

    ws2['A1'] = "PROJEÇÃO INTELIGENTE (PRÓXIMOS MESES)"
    ws2['A1'].font = Font(name='Calibri', bold=True, size=14, color='1E1B4B')

    headers_proj = ["Mês", "Gastos Previstos", "Receitas Previstas", "Sobra"]
    for col, text in enumerate(headers_proj, start=1):
        c = ws2.cell(row=3, column=col, value=text)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill(start_color='8B5CF6', end_color='8B5CF6', fill_type='solid')
        c.alignment = Alignment(horizontal='center')

    for i, p in enumerate(projecao_dados, start=4):
        ws2.cell(row=i, column=1, value=p['nome']).alignment = Alignment(horizontal='center')
        ws2.cell(row=i, column=2, value=float(p['gasto'])).number_format = 'R$ #,##0.00'
        ws2.cell(row=i, column=3, value=float(p['receita'])).number_format = 'R$ #,##0.00'
        ws2.cell(row=i, column=4, value=float(p['sobra'])).number_format = 'R$ #,##0.00'

    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "Fluxo de Caixa Futuro"
    data_bar = Reference(ws2, min_col=2, max_col=3, min_row=3, max_row=3 + len(projecao_dados))
    cats_bar = Reference(ws2, min_col=1, min_row=4, max_row=3 + len(projecao_dados))
    bar.add_data(data_bar, titles_from_data=True)
    bar.set_categories(cats_bar)
    ws2.add_chart(bar, "F3")

    ws2['A15'] = "💡 DICAS DE ECONOMIA E MELHORIA"
    ws2['A15'].font = Font(bold=True, size=12, color='DC2626')

    dicas_texto = [
        "1. Analise as categorias do gráfico de rosca. Onde há maior fatia?",
        "2. Tente reduzir em 10% o maior gasto não essencial deste mês.",
        "3. Compare a coluna 'Receitas' com 'Gastos' na projeção acima.",
        "4. Se a projeção mostrar meses no negativo, reveja compras parceladas.",
        "5. Não se esqueça de alimentar a Reserva de Emergência no sistema."
    ]
    for i, dica in enumerate(dicas_texto, start=16):
        ws2.cell(row=i, column=1, value=dica).font = Font(italic=True, color='475569')

    # ABA 3: EXTRATO DETALHADO
    ws3 = wb.create_sheet(title="Extrato Detalhado")

    ws3['A1'] = f"EXTRATO DE GASTOS - {mes_ativo}"
    ws3['A1'].font = Font(bold=True, size=14, color='1E1B4B')
    ws3.merge_cells('A1:E1')

    headers_ext = ['Data', 'Descrição', 'Pessoa', 'Categoria', 'Valor']
    for col, text in enumerate(headers_ext, start=1):
        c = ws3.cell(row=3, column=col, value=text)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill(start_color='EC4899', end_color='EC4899', fill_type='solid')
        c.alignment = Alignment(horizontal='center')

    for i, g in enumerate(gastos, start=4):
        ws3.cell(row=i, column=1, value=g.data.strftime('%d/%m/%Y'))
        ws3.cell(row=i, column=2, value=g.descricao)
        ws3.cell(row=i, column=3, value=g.pessoa.nome)
        ws3.cell(row=i, column=4, value=g.categoria.nome)
        c_val = ws3.cell(row=i, column=5, value=float(g.valor))
        c_val.number_format = 'R$ #,##0.00'
        c_val.font = Font(color='DC2626', bold=True)

    linha_total = 4 + len(gastos)
    ws3.cell(row=linha_total, column=4, value="TOTAL:").font = Font(bold=True)
    c_tot = ws3.cell(row=linha_total, column=5, value=float(total_gasto_mes))
    c_tot.number_format = 'R$ #,##0.00'
    c_tot.font = Font(bold=True, color='DC2626')

    # Ajustar colunas (com proteção para células mescladas)
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            cell_inicial = col[0]
            if not hasattr(cell_inicial, 'column_letter'):
                continue

            max_length = 0
            column_letter = cell_inicial.column_letter

            for cell in col:
                try:
                    if hasattr(cell, 'value') and cell.value is not None:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except Exception:
                    pass

            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Relatorio_Profissional_{mes_ativo.ano}_{mes_ativo.mes}.xlsx"'

    wb.save(response)
    return response


# ==========================================
# TELA PRINCIPAL (AGENDA / CONTAS A VENCER)
# ==========================================
@login_required
def home(request):
    hoje = date.today()

    ano_get = request.GET.get('ano')
    mes_get = request.GET.get('mes')

    if ano_get and mes_get:
        ano_base = int(ano_get)
        mes_base = int(mes_get)
    else:
        mes_exibido = Mes.objects.filter(ano=hoje.year, mes=hoje.month, fechado=False).first()

        if not mes_exibido:
            mes_exibido = Mes.objects.filter(fechado=False).order_by('ano', 'mes').filter(
                models.Q(ano__gt=hoje.year) | models.Q(ano=hoje.year, mes__gte=hoje.month)
            ).first()

        if not mes_exibido:
            mes_exibido = Mes.objects.filter(fechado=False).order_by('-ano', '-mes').first()

        if not mes_exibido:
            mes_exibido, _ = Mes.objects.get_or_create(ano=hoje.year, mes=hoje.month)

        ano_base = mes_exibido.ano
        mes_base = mes_exibido.mes

    mes_exibido = Mes.objects.filter(ano=ano_base, mes=mes_base).first()
    if not mes_exibido:
        mes_exibido = Mes.objects.filter(id=request.session.get('mes_ativo_id')).first()
        if not mes_exibido:
            mes_exibido, _ = Mes.objects.get_or_create(ano=hoje.year, mes=hoje.month)
        ano_base = mes_exibido.ano
        mes_base = mes_exibido.mes

    request.session['mes_ativo_id'] = mes_exibido.id

    data_atual = date(ano_base, mes_base, 1)
    data_anterior = data_atual - relativedelta(months=1)
    data_posterior = data_atual + relativedelta(months=1)

    contas_individuais = Gasto.objects.filter(
        cartao__isnull=True,
        data__year=ano_base,
        data__month=mes_base
    ).order_by('data')

    gastos_cartao = Gasto.objects.filter(
        cartao__isnull=False,
        mes=mes_exibido
    )

    cartoes_resumo = {}
    for g in gastos_cartao:
        if g.cartao.id not in cartoes_resumo:
            dia_venc = g.cartao.dia_vencimento if g.cartao.dia_vencimento else 1
            data_vencimento = date(ano_base, mes_base, dia_venc)

            cartoes_resumo[g.cartao.id] = {
                'cartao': g.cartao,
                'total': Decimal('0'),
                'vencimento': data_vencimento,
            }
        cartoes_resumo[g.cartao.id]['total'] += g.valor

    total_geral_gastos = sum(g.valor for g in contas_individuais) + sum(c['total'] for c in cartoes_resumo.values())
    total_geral_receitas = Receita.objects.filter(mes=mes_exibido).aggregate(total=Sum('valor'))['total'] or Decimal('0')
    saldo_geral = total_geral_receitas - total_geral_gastos

    if ano_base == hoje.year and mes_base == hoje.month:
        proximas = [c for c in contas_individuais if c.data >= hoje]
        atrasadas = [c for c in contas_individuais if c.data < hoje]
    else:
        proximas = contas_individuais
        atrasadas = []

    lista_proximas = proximas

    todos_gastos_mes = list(contas_individuais) + list(gastos_cartao)
    todos_gastos_mes = sorted(todos_gastos_mes, key=lambda x: x.data, reverse=True)

    eventos_calendario = []
    for g in contas_individuais:
        cor = '#8b5cf6' if g.categoria.tipo == 'E' else '#ec4899'
        eventos_calendario.append({
            'title': f"R$ {g.valor} - {g.descricao}",
            'start': g.data.isoformat(),
            'backgroundColor': cor,
            'borderColor': '#222222',
            'textColor': '#ffffff'
        })

    for c_id, c_info in cartoes_resumo.items():
        cor_hex = c_info['cartao'].cor if c_info['cartao'].cor else '#d1b3e2'
        eventos_calendario.append({
            'title': f"Fatura {c_info['cartao'].nome}: R$ {c_info['total']}",
            'start': c_info['vencimento'].isoformat(),
            'backgroundColor': cor_hex,
            'borderColor': '#222222',
            'textColor': '#111111'
        })

    eventos_json = json.dumps(eventos_calendario)

    form_gasto = GastoForm(initial={'data': date(ano_base, mes_base, 1)})

    return render(request, 'financas/home.html', {
        'mes_ativo': mes_exibido,
        'hoje': hoje,
        'cartoes_resumo': list(cartoes_resumo.values()),
        'lista_proximas': lista_proximas,
        'atrasadas': atrasadas,
        'total_atrasado': sum(g.valor for g in atrasadas),
        'eventos_json': eventos_json,
        'ano_anterior': data_anterior.year,
        'mes_anterior': data_anterior.month,
        'ano_posterior': data_posterior.year,
        'mes_posterior': data_posterior.month,
        'form_gasto': form_gasto,
        'total_geral_gastos': total_geral_gastos,
        'total_geral_receitas': total_geral_receitas,
        'saldo_geral': saldo_geral,
        'lista_extrato': todos_gastos_mes,
    })


# ==========================================
# PERFIL
# ==========================================
@login_required
def perfil(request):
    perfil, created = Perfil.objects.get_or_create(usuario=request.user)

    if request.method == 'POST':
        if 'salvar_foto' in request.POST:
            form_perfil = PerfilForm(request.POST, request.FILES, instance=perfil)
            form_senha = AlterarSenhaForm(user=request.user)

            if form_perfil.is_valid():
                form_perfil.save()
                messages.success(request, 'Foto de perfil atualizada!')
                return redirect('perfil')
            else:
                messages.error(request, 'Erro ao salvar a foto. Verifique o arquivo.')

        elif 'salvar_senha' in request.POST:
            form_perfil = PerfilForm(instance=perfil)
            form_senha = AlterarSenhaForm(user=request.user, data=request.POST)

            if form_senha.is_valid():
                user = form_senha.save()
                update_session_auth_hash(request, user)
                messages.success(request, 'Sua senha foi alterada com sucesso!')
                return redirect('perfil')
            else:
                messages.error(request, 'Erro ao alterar a senha. Verifique os dados.')
    else:
        form_perfil = PerfilForm(instance=perfil)
        form_senha = AlterarSenhaForm(user=request.user)

    return render(request, 'financas/perfil.html', {
        'perfil': perfil,
        'form_perfil': form_perfil,
        'form_senha': form_senha,
    })